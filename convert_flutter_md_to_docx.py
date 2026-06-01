#!/usr/bin/env python3
"""
Convert flutter-interview-middle.md to a nicely formatted .docx Word file.
"""

import re
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ─── Color constants ────────────────────────────────────────────────────────
DARK_BLUE   = RGBColor(0x00, 0x33, 0x66)
DARK_RED    = RGBColor(0xCC, 0x00, 0x00)
GREEN       = RGBColor(0,    128,  0)
GRAY_CODE   = RGBColor(0xF0, 0xF0, 0xF0)
INLINE_GRAY = RGBColor(0xE8, 0xE8, 0xE8)

# ─── Terms to explain (first occurrence only) ───────────────────────────────
FLUTTER_TERMS = {
    "Null Safety":        "Tính năng của Dart đảm bảo biến không thể là null trừ khi được khai báo tường minh với `?`, giúp loại bỏ lỗi null pointer runtime.",
    "Future":             "Đại diện cho một giá trị chưa có ngay, sẽ hoàn thành trong tương lai (một lần duy nhất) — tương tự Promise trong JavaScript.",
    "Stream":             "Luồng dữ liệu bất đồng bộ phát ra nhiều giá trị theo thời gian, như dòng sự kiện liên tục.",
    "Isolate":            "Luồng thực thi độc lập trong Dart với bộ nhớ riêng, không chia sẻ memory với isolate khác, giao tiếp qua message passing.",
    "Mixin":              "Cơ chế tái sử dụng code trong Dart: thêm tính năng vào class mà không cần kế thừa, tránh vấn đề đa kế thừa.",
    "Widget Tree":        "Cây phân cấp của tất cả các widget trong app Flutter, là cấu trúc mô tả giao diện người dùng.",
    "Element Tree":       "Cây nội bộ Flutter quản lý vòng đời widget, là cầu nối giữa Widget Tree và Render Tree.",
    "Render Tree":        "Cây chứa các đối tượng thực sự đo lường, layout và vẽ lên màn hình.",
    "BuildContext":       "Vị trí của widget trong Widget Tree, dùng để truy cập dữ liệu từ widget cha (InheritedWidget, Theme...).",
    "InheritedWidget":    "Widget đặc biệt cho phép truyền dữ liệu xuống toàn bộ cây con mà không cần truyền qua constructor từng cấp.",
    "BLoC Pattern":       "Business Logic Component — kiến trúc tách biệt UI và business logic, dùng Stream để giao tiếp.",
    "Cubit":              "Phiên bản đơn giản hơn của Bloc, thay vì xử lý Event thì gọi trực tiếp method để emit State.",
    "Riverpod":           "Thư viện quản lý state cải tiến từ Provider, an toàn hơn về compile-time, không phụ thuộc vào BuildContext.",
    "Provider":           "Gói quản lý state phổ biến trong Flutter, bọc InheritedWidget để dễ sử dụng hơn.",
    "Navigator 2.0":      "API điều hướng declarative mới của Flutter, kiểm soát navigation stack bằng state thay vì gọi trực tiếp.",
    "Deep Linking":       "Khả năng mở thẳng màn hình cụ thể trong app từ URL bên ngoài (web, thông báo...).",
    "Jank":               "Hiện tượng app bị giật, xảy ra khi frame mất hơn 16ms để render (không đạt 60fps).",
    "RepaintBoundary":    "Widget tạo layer riêng biệt ngăn việc vẽ lại lan sang widget xung quanh khi chỉ một phần thay đổi.",
    "Platform Channel":   "Cơ chế giao tiếp giữa Flutter (Dart) và code native (Kotlin/Swift) để gọi API native.",
    "Skia":               "Engine đồ họa 2D cũ của Flutter, vẽ toàn bộ UI từ đầu mỗi frame.",
    "Impeller":           "Engine render mới của Flutter (thay thế Skia), pre-compile shader để loại bỏ jank khi chạy lần đầu.",
    "Clean Architecture": "Kiến trúc chia app thành 3 tầng: Presentation, Domain, Data — tách biệt hoàn toàn business logic khỏi UI và data source.",
    "Dependency Injection":"Kỹ thuật cung cấp dependency (đối tượng phụ thuộc) từ bên ngoài vào class thay vì tạo trong class, giúp dễ test và thay thế.",
    "GetIt":              "Thư viện service locator phổ biến trong Flutter, dùng để đăng ký và lấy dependency toàn cục.",
    "compute()":          "Hàm tiện ích của Flutter chạy một function nặng trên Isolate riêng, tránh block UI thread.",
    "const constructor":  "Constructor đặc biệt trong Dart tạo ra đối tượng compile-time constant, Flutter tái sử dụng thay vì tạo mới → tối ưu hiệu năng.",
}

explained_terms = set()


# ─── XML helpers ────────────────────────────────────────────────────────────
def set_cell_bg(cell, hex_color):
    """Set background color of a table cell."""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    tcPr.append(shd)


def add_top_border(paragraph):
    """Add a top border to a paragraph (used for H2 headings)."""
    pPr  = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    top  = OxmlElement('w:top')
    top.set(qn('w:val'),   'single')
    top.set(qn('w:sz'),    '12')
    top.set(qn('w:space'), '4')
    top.set(qn('w:color'), '003366')
    pBdr.append(top)
    pPr.append(pBdr)


def shade_paragraph_bg(paragraph, hex_color):
    """Add shading to a paragraph background."""
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    pPr.append(shd)


def set_para_indent(paragraph, left_cm):
    """Set left indent on a paragraph."""
    pPr  = paragraph._p.get_or_add_pPr()
    ind  = OxmlElement('w:ind')
    twips = int(left_cm * 567)          # 1 cm ≈ 567 twips
    ind.set(qn('w:left'), str(twips))
    pPr.append(ind)


# ─── Inline-markup helpers ───────────────────────────────────────────────────
def add_inline_run(paragraph, text, bold=False, italic=False,
                   color=None, font_name=None, font_size=None,
                   bg_color_rgb=None):
    """Append a run with mixed formatting to a paragraph."""
    run = paragraph.add_run(text)
    run.bold   = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    if font_name:
        run.font.name = font_name
    if font_size:
        run.font.size = Pt(font_size)
    if bg_color_rgb:
        rPr = run._r.get_or_add_rPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'),   'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'),  bg_color_rgb)
        rPr.append(shd)
    return run


def parse_inline(paragraph, text, base_font_size=11, base_bold=False,
                 base_color=None, base_italic=False):
    """
    Parse inline markdown in `text` and add formatted runs to `paragraph`.
    Handles: **bold**, `inline code`, mixed combos.
    """
    # Pattern: **bold** | `code`
    pattern = re.compile(r'(\*\*(.+?)\*\*|`([^`]+)`)')
    last = 0
    for m in pattern.finditer(text):
        # plain text before this match
        if m.start() > last:
            chunk = text[last:m.start()]
            add_inline_run(paragraph, chunk,
                           bold=base_bold, italic=base_italic,
                           color=base_color, font_size=base_font_size)
        if m.group().startswith('**'):
            inner = m.group(2)
            add_inline_run(paragraph, inner, bold=True, italic=base_italic,
                           color=base_color or DARK_BLUE,
                           font_size=base_font_size)
        else:   # backtick code
            inner = m.group(3)
            add_inline_run(paragraph, inner, font_name='Courier New',
                           font_size=9, bg_color_rgb='E8E8E8')
        last = m.end()
    # remainder
    if last < len(text):
        add_inline_run(paragraph, text[last:],
                       bold=base_bold, italic=base_italic,
                       color=base_color, font_size=base_font_size)


# ─── Term-explanation helper ─────────────────────────────────────────────────
def maybe_add_explanation(doc, heading_text):
    """
    If the heading contains an unexlained Flutter term, insert a 💡 note.
    """
    global explained_terms
    for term, explanation in FLUTTER_TERMS.items():
        # match term in heading (case-sensitive, whole-word-ish)
        if term in heading_text and term not in explained_terms:
            explained_terms.add(term)
            p = doc.add_paragraph()
            set_para_indent(p, 0.5)
            run = p.add_run(f"💡 Giải thích: {explanation}")
            run.font.size   = Pt(10)
            run.font.name   = 'Calibri'
            run.italic      = True
            run.font.color.rgb = GREEN
            return
    # Also scan term list body text for first occurrences (done at line level)


def check_body_terms(doc, line_text):
    """Insert 💡 note if a Flutter term appears for the first time in body text."""
    global explained_terms
    for term, explanation in FLUTTER_TERMS.items():
        if term in line_text and term not in explained_terms:
            explained_terms.add(term)
            p = doc.add_paragraph()
            set_para_indent(p, 0.5)
            run = p.add_run(f"💡 Giải thích: {explanation}")
            run.font.size      = Pt(10)
            run.font.name      = 'Calibri'
            run.italic         = True
            run.font.color.rgb = GREEN
            return p
    return None


# ─── Table helpers ───────────────────────────────────────────────────────────
def add_markdown_table(doc, table_lines):
    """Render a markdown pipe-table as a Word table."""
    rows = []
    for line in table_lines:
        line = line.strip()
        if not line or re.match(r'^\|?[-| :]+\|?$', line):
            continue
        cells = [c.strip() for c in re.split(r'(?<!\\)\|', line) if c.strip()]
        rows.append(cells)

    if not rows:
        return

    col_count = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=col_count)
    tbl.style = 'Table Grid'

    for r_idx, row_cells in enumerate(rows):
        row = tbl.rows[r_idx]
        for c_idx, cell_text in enumerate(row_cells):
            if c_idx >= col_count:
                break
            cell = row.cells[c_idx]
            cell.text = ''
            p = cell.paragraphs[0]
            is_header = (r_idx == 0)
            parse_inline(p, cell_text, base_font_size=9,
                         base_bold=is_header,
                         base_color=DARK_BLUE if is_header else None)
            if is_header:
                set_cell_bg(cell, 'D0D8E8')
            for run in p.runs:
                run.font.size = Pt(9)

    # Fill empty cells
    for r_idx, row_cells in enumerate(rows):
        row = tbl.rows[r_idx]
        for c_idx in range(len(row_cells), col_count):
            row.cells[c_idx].text = ''


# ─── Code block helper ───────────────────────────────────────────────────────
def add_code_block(doc, code_lines):
    """Render a code block with Courier New, gray background, indented."""
    for line in code_lines:
        p = doc.add_paragraph()
        p.style = doc.styles['Normal']
        set_para_indent(p, 0.5)
        shade_paragraph_bg(p, 'F0F0F0')
        # preserve leading spaces
        run = p.add_run(line if line else ' ')
        run.font.name = 'Courier New'
        run.font.size = Pt(9)
        # reduce space after
        p.paragraph_format.space_after  = Pt(0)
        p.paragraph_format.space_before = Pt(0)


# ─── Main converter ──────────────────────────────────────────────────────────
def convert(md_path, docx_path):
    with open(md_path, encoding='utf-8') as f:
        raw = f.read()

    lines = raw.splitlines()

    doc = Document()

    # ── Page margins ──────────────────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin    = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin   = Cm(2)
        section.right_margin  = Cm(2)

    # ── Default paragraph style ───────────────────────────────────────────────
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    i = 0
    n = len(lines)

    while i < n:
        line = lines[i]

        # ── Skip horizontal rules ─────────────────────────────────────────────
        if re.match(r'^-{3,}$', line.strip()):
            i += 1
            continue

        # ── Code block ────────────────────────────────────────────────────────
        if line.strip().startswith('```'):
            i += 1
            code_lines = []
            while i < n and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            i += 1  # skip closing ```
            add_code_block(doc, code_lines)
            continue

        # ── Markdown table ────────────────────────────────────────────────────
        if re.match(r'^\s*\|', line):
            table_lines = []
            while i < n and re.match(r'^\s*\|', lines[i]):
                table_lines.append(lines[i])
                i += 1
            add_markdown_table(doc, table_lines)
            continue

        # ── H1 (Title) ────────────────────────────────────────────────────────
        m = re.match(r'^# (.+)', line)
        if m:
            text = m.group(1).strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after  = Pt(12)
            run = p.add_run(text)
            run.bold            = True
            run.font.size       = Pt(20)
            run.font.color.rgb  = DARK_BLUE
            run.font.name       = 'Calibri'
            i += 1
            continue

        # ── H2 (PHẦN X) ──────────────────────────────────────────────────────
        m = re.match(r'^## (.+)', line)
        if m:
            text = m.group(1).strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after  = Pt(4)
            add_top_border(p)
            run = p.add_run(text)
            run.bold           = True
            run.font.size      = Pt(16)
            run.font.color.rgb = DARK_BLUE
            run.font.name      = 'Calibri'
            # Check terms in H2 text
            for term in FLUTTER_TERMS:
                if term in text and term not in explained_terms:
                    explained_terms.add(term)
                    ep = doc.add_paragraph()
                    set_para_indent(ep, 0.5)
                    er = ep.add_run(f"💡 Giải thích: {FLUTTER_TERMS[term]}")
                    er.font.size      = Pt(10)
                    er.font.name      = 'Calibri'
                    er.italic         = True
                    er.font.color.rgb = GREEN
                    break
            i += 1
            continue

        # ── H3 ────────────────────────────────────────────────────────────────
        m = re.match(r'^### (.+)', line)
        if m:
            text = m.group(1).strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after  = Pt(4)
            set_para_indent(p, 0.3)
            run = p.add_run(text)
            run.bold           = True
            run.font.size      = Pt(13)
            run.font.color.rgb = DARK_BLUE
            run.font.name      = 'Calibri'
            # check terms
            for term in FLUTTER_TERMS:
                if term in text and term not in explained_terms:
                    explained_terms.add(term)
                    ep = doc.add_paragraph()
                    set_para_indent(ep, 0.5)
                    er = ep.add_run(f"💡 Giải thích: {FLUTTER_TERMS[term]}")
                    er.font.size      = Pt(10)
                    er.font.name      = 'Calibri'
                    er.italic         = True
                    er.font.color.rgb = GREEN
                    break
            i += 1
            continue

        # ── H4 ────────────────────────────────────────────────────────────────
        m = re.match(r'^#### (.+)', line)
        if m:
            text = m.group(1).strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after  = Pt(2)
            run = p.add_run(text)
            run.bold      = True
            run.font.size = Pt(12)
            run.font.name = 'Calibri'
            i += 1
            continue

        # ── Bullet list ───────────────────────────────────────────────────────
        m = re.match(r'^(\s*)[-*] (.+)', line)
        if m:
            text = m.group(2).strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            # Is it a Q: question?
            if text.startswith('Q:'):
                add_inline_run(p, text, bold=True, color=DARK_RED, font_size=11)
            else:
                parse_inline(p, text, base_font_size=11)
            check_body_terms(doc, text)
            i += 1
            continue

        # ── Numbered list ─────────────────────────────────────────────────────
        m = re.match(r'^(\s*)\d+\. (.+)', line)
        if m:
            text = m.group(2).strip()
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(2)
            parse_inline(p, text, base_font_size=11)
            check_body_terms(doc, text)
            i += 1
            continue

        # ── Checklist (- [ ]) ─────────────────────────────────────────────────
        m = re.match(r'^- \[[ x]\] (.+)', line)
        if m:
            text = m.group(1).strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            checked = '☑ ' if '[x]' in line else '☐ '
            parse_inline(p, checked + text, base_font_size=11)
            i += 1
            continue

        # ── Empty line ────────────────────────────────────────────────────────
        if not line.strip():
            i += 1
            continue

        # ── Normal paragraph ──────────────────────────────────────────────────
        text = line.strip()

        # Q: question lines
        if text.startswith('**Q:') or re.match(r'^\*\*Q\d*:', text):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after  = Pt(4)
            # strip outer **
            inner = text.strip('*').strip()
            add_inline_run(p, inner, bold=True, color=DARK_RED, font_size=11)
            check_body_terms(doc, inner)
            i += 1
            continue

        # standalone Q: lines without **
        if text.startswith('Q:'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after  = Pt(4)
            add_inline_run(p, text, bold=True, color=DARK_RED, font_size=11)
            check_body_terms(doc, text)
            i += 1
            continue

        # Otherwise plain text with inline markup
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        parse_inline(p, text, base_font_size=11)
        check_body_terms(doc, text)
        i += 1

    doc.save(docx_path)
    return docx_path


# ─── Entry point ─────────────────────────────────────────────────────────────
if __name__ == '__main__':
    MD_PATH   = r'c:\Users\ALLIANCEITSC\Documents\learning\flutter-interview-middle.md'
    DOCX_PATH = r'c:\Users\ALLIANCEITSC\Documents\learning\flutter-interview-middle.docx'

    print(f"Converting: {MD_PATH}")
    out = convert(MD_PATH, DOCX_PATH)

    if os.path.exists(out):
        size = os.path.getsize(out)
        print(f"✅ Created: {out}")
        print(f"📄 File size: {size:,} bytes ({size/1024:.1f} KB)")
    else:
        print("❌ File was NOT created!")
